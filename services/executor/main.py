import asyncio
import glob
import json
import logging
import os
import shutil
import traceback
import uuid
from contextlib import asynccontextmanager
from datetime import datetime, timezone

import httpx
import openpyxl
from fastapi import BackgroundTasks, FastAPI, File, UploadFile
from fastapi.responses import StreamingResponse

from config import DASHBOARD_URL, JAVA_DIR, PORT, SHARED_DIR, resolve_java_dir
from excel_writer import write as excel_write
from mvn_runner import run as mvn_run
from surefire_parser import parse as surefire_parse
from testng_generator import generate as testng_gen
from js_runner import run as js_run
from ts_runner import run as ts_run
from py_runner import run as py_run

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(name)s %(levelname)s %(message)s")
logger = logging.getLogger("executor-service")

_jobs: dict = {}
_active_runs: int = 0
_run_queue: asyncio.Queue = asyncio.Queue()


@asynccontextmanager
async def lifespan(app: FastAPI):
    logger.info("=" * 50)
    logger.info("Executor service online — ready for test runs")
    logger.info("Deployment type: ALWAYS-ON")
    logger.info("Handling: user story tests + regression runs")
    logger.info("=" * 50)
    yield
    logger.info("Executor service shutting down")


app = FastAPI(
    title="Executor Service",
    description="Always-on test executor. Handles regression + user story runs.",
    version="1.0.0",
    lifespan=lifespan
)


async def _notify(app_id: str, stage: str, message: str, level: str = "INFO"):
    try:
        async with httpx.AsyncClient(timeout=5) as client:
            await client.post(f"{DASHBOARD_URL}/api/events", json={
                "app_id": app_id, "stage": stage, "message": message, "level": level
            })
    except Exception:
        pass


async def _execute(run_id: str, app_id: str, java_dir: str,
                   selected_tests: list[str], test_data: list[dict],
                   framework: str = "java", app_url: str = ""):
    global _active_runs
    _active_runs += 1
    _jobs[run_id]["status"] = "running"
    testdata_path = None
    testng_path = None

    try:
        n = len(selected_tests)
        logger.info(f"Executing run {run_id} — {n} test cases selected [{framework}]")
        await _notify(app_id, "EXECUTE", f"Run started — {n} tests [{framework}]")

        if framework in ("js", "ts", "python"):
            runner = {"js": js_run, "ts": ts_run, "python": py_run}[framework]
            fw_result = await runner(java_dir, run_id, app_id, selected_tests, test_data,
                                     DASHBOARD_URL, app_url)
            exit_code = fw_result.get("exit_code", 1)
            results = fw_result.get("results", [])
        else:
            # Java / Maven path
            testdata_path = excel_write(test_data, java_dir)
            testng_path = testng_gen(selected_tests, java_dir, run_id, app_id)
            mvn_result = await mvn_run(java_dir, run_id, app_id, DASHBOARD_URL)
            exit_code = mvn_result.get("exit_code", 1)
            results = surefire_parse(java_dir, run_id)

        passed = sum(1 for r in results if r["status"] == "PASSED")
        failed = sum(1 for r in results if r["status"] == "FAILED")

        async with httpx.AsyncClient(timeout=10) as client:
            await client.post(f"{DASHBOARD_URL}/api/execute/results", json={
                "run_id": run_id,
                "results": results,
                "mvn_exit_code": exit_code
            })

        _jobs[run_id].update({
            "status": "done",
            "passed": passed,
            "failed": failed,
            "skipped": len(results) - passed - failed,
            "mvn_exit_code": exit_code
        })

        logger.info(f"Run {run_id} complete — {passed} passed, {failed} failed")
        await _notify(app_id, "EXECUTE", f"Run {run_id} complete — {passed} passed, {failed} failed")

    except Exception:
        tb = traceback.format_exc()
        logger.error(f"Executor run {run_id} failed:\n{tb}")
        await _notify(app_id, "EXECUTE", f"Run {run_id} failed: {tb[:300]}", "ERROR")
        _jobs[run_id]["status"] = "error"
    finally:
        _active_runs -= 1
        for path in [testdata_path, testng_path]:
            if path and os.path.exists(path):
                try:
                    os.remove(path)
                except Exception:
                    pass


@app.post("/execute/run")
async def execute_run(body: dict, background_tasks: BackgroundTasks):
    app_id         = body["app_id"]
    run_id         = body.get("run_id", "run-" + uuid.uuid4().hex[:8])
    framework      = body.get("framework", "java")
    app_url        = body.get("app_url", "")
    # accept framework-specific dir keys or fall back to java_dir
    dir_key_map    = {"js": "js_dir", "ts": "ts_dir", "python": "py_dir"}
    dir_key        = dir_key_map.get(framework, "java_dir")
    raw_dir        = body.get(dir_key) or body.get("java_dir", JAVA_DIR)
    fw_dir         = resolve_java_dir(raw_dir)
    selected_tests = body.get("selected_tests", [])
    test_data      = body.get("test_data", [])

    _jobs[run_id] = {"status": "started", "passed": 0, "failed": 0, "skipped": 0, "framework": framework}
    background_tasks.add_task(_execute, run_id, app_id, fw_dir, selected_tests, test_data, framework, app_url)
    return {"run_id": run_id, "status": "STARTED"}


@app.get("/status/{run_id}")
async def status(run_id: str):
    job = _jobs.get(run_id, {"status": "not_found"})
    return {"run_id": run_id, **job}


@app.get("/health")
async def health():
    return {
        "service": "executor",
        "status": "running" if _active_runs > 0 else "idle",
        "deployment_type": "always-on",
        "active_runs": _active_runs,
        "port": PORT
    }


# ── helpers ──────────────────────────────────────────────────────────────────

def _tc_json_path(app_id: str) -> str:
    return os.path.join(SHARED_DIR, "outputs", app_id, "test_cases.json")


def _load_tc_json(app_id: str) -> dict:
    path = _tc_json_path(app_id)
    if os.path.exists(path):
        with open(path) as f:
            return json.load(f)
    return {}


def _results_dir(app_id: str) -> str:
    d = os.path.join(SHARED_DIR, "outputs", app_id, "results")
    os.makedirs(d, exist_ok=True)
    return d


# ── export template ───────────────────────────────────────────────────────────

@app.post("/execute/export-template")
async def export_template(body: dict):
    """Build an Excel template (one TestData sheet) for the selected test methods."""
    app_id       = body["app_id"]
    method_names = body.get("method_names", [])

    tc_map = _load_tc_json(app_id)

    # collect ordered unique param names for the requested methods
    seen_params: list[str] = []
    rows: list[str] = []
    for method in method_names:
        entry = next(
            (v for v in tc_map.values() if v.get("method_name") == method),
            None,
        )
        if entry is None:
            continue
        rows.append(method)
        for p in entry.get("parameters", []):
            name = p.get("name", "")
            if name and name not in seen_params:
                seen_params.append(name)

    wb = openpyxl.Workbook()
    ws = wb.create_sheet("TestData")
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]
    ws.append(["Test Case Name"] + seen_params)
    for method in rows:
        ws.append([method] + [""] * len(seen_params))

    import io
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{app_id}-test-data-template.xlsx"'},
    )


# ── import data ───────────────────────────────────────────────────────────────

@app.post("/execute/import-data")
async def import_data(app_id: str, file: UploadFile = File(...), java_dir: str | None = None):
    """Import a filled test-data Excel, write per-method JSON files, regen testng.xml."""
    java_dir = resolve_java_dir(java_dir or os.environ.get("JAVA_DIR", JAVA_DIR))
    tc_map   = _load_tc_json(app_id)

    contents = await file.read()
    import io
    wb = openpyxl.load_workbook(io.BytesIO(contents))
    ws = wb["TestData"]

    headers = [str(c.value or "").strip() for c in ws[1]]
    if "Test Case Name" not in headers:
        return {"error": "Missing 'Test Case Name' column"}

    name_idx = headers.index("Test Case Name")

    # group rows by method name
    grouped: dict[str, list[dict]] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        method = str(row[name_idx] or "").strip()
        if not method:
            continue
        obj = {headers[i]: str(v or "") for i, v in enumerate(row) if i != name_idx and headers[i]}
        grouped.setdefault(method, []).append(obj)

    # wipe testdata folder
    testdata_dir = os.path.join(java_dir, "src", "test", "resources", "testdata")
    if os.path.exists(testdata_dir):
        shutil.rmtree(testdata_dir)
    os.makedirs(testdata_dir, exist_ok=True)

    imported: dict[str, int] = {}
    for method, data_rows in grouped.items():
        entry = next(
            (v for v in tc_map.values() if v.get("method_name") == method),
            None,
        )
        param_names = [p["name"] for p in (entry or {}).get("parameters", [])] if entry else list(data_rows[0].keys())
        filtered = [{p: r.get(p, "") for p in param_names} for r in data_rows]
        out = os.path.join(testdata_dir, f"{method}.json")
        with open(out, "w") as f:
            json.dump(filtered, f, indent=2)
        imported[method] = len(filtered)

    # regenerate testng.xml with imported methods
    testng_path = testng_gen(list(grouped.keys()), java_dir, "suite", app_id)

    return {"imported": imported, "testng_xml": testng_path}


# ── run suite (SSE) ───────────────────────────────────────────────────────────

@app.post("/execute/run-suite")
async def run_suite(body: dict):
    """Run the pre-generated testng.xml and stream Maven output via SSE."""
    app_id   = body["app_id"]
    java_dir = resolve_java_dir(body.get("java_dir", JAVA_DIR))
    app_url  = body.get("app_url", "")
    run_id   = "run-" + uuid.uuid4().hex[:8]
    started  = datetime.now(timezone.utc).isoformat()

    async def _stream():
        testng_file = os.path.join(java_dir, "testng.xml")
        mvn_args = [
            "mvn", "test",
            f"-Dsurefire.suiteXmlFiles={testng_file}",
        ]
        if app_url:
            mvn_args.append(f"-Dapp.url={app_url}")
        proc = await asyncio.create_subprocess_exec(
            *mvn_args,
            cwd=java_dir,
            stdout=asyncio.subprocess.PIPE,
            stderr=asyncio.subprocess.STDOUT,
        )

        async for raw in proc.stdout:  # type: ignore[union-attr]
            line = raw.decode(errors="replace").rstrip()
            yield f"data: {json.dumps({'line': line})}\n\n"

        await proc.wait()
        exit_code = proc.returncode

        # parse surefire results
        results = surefire_parse(java_dir, run_id)

        # write result file
        result_obj = {
            "run_id":      run_id,
            "app_id":      app_id,
            "started_at":  started,
            "finished_at": datetime.now(timezone.utc).isoformat(),
            "exit_code":   exit_code,
            "tests": [
                {
                    "method":      r["tc_name"],
                    "status":      r["status"],
                    "duration_ms": r["duration_ms"],
                    "error":       r.get("failure_msg", ""),
                }
                for r in results
            ],
        }
        out_path = os.path.join(_results_dir(app_id), f"{run_id}.json")
        with open(out_path, "w") as f:
            json.dump(result_obj, f, indent=2)

        yield f"data: {json.dumps({'exit_code': exit_code, 'result_file': f'results/{run_id}.json', 'tests': result_obj['tests']})}\n\n"

    return StreamingResponse(_stream(), media_type="text/event-stream")


# ── results ───────────────────────────────────────────────────────────────────

@app.get("/execute/results")
async def list_results(app_id: str, limit: int = 20):
    """List execution results for an app, newest first."""
    results_dir = _results_dir(app_id)
    files = glob.glob(os.path.join(results_dir, "*.json"))
    out = []
    for fpath in files:
        try:
            with open(fpath) as f:
                out.append(json.load(f))
        except Exception:
            pass
    out.sort(key=lambda r: r.get("started_at", ""), reverse=True)
    return {"results": out[:limit]}


@app.get("/execute/results/{run_id}")
async def get_result(run_id: str, app_id: str):
    """Get a single execution result by run_id."""
    fpath = os.path.join(_results_dir(app_id), f"{run_id}.json")
    if not os.path.exists(fpath):
        return {"error": f"Result not found: {run_id}"}
    with open(fpath) as f:
        return json.load(f)
